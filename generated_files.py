"""Bounded extraction of user-supplied generated output files."""

from __future__ import annotations

import hashlib
import io
import zipfile
from pathlib import Path

MAX_FILE_BYTES = 10 * 1024 * 1024
MAX_TEXT_CHARS = 30_000
FORMATS = {".html": "html", ".htm": "html", ".pdf": "pdf", ".xlsx": "xlsx"}


def inspect_generated_file(filename: str, data: bytes) -> dict:
    if not isinstance(filename, str) or not filename or len(filename) > 255 or any(ord(c) < 32 for c in filename):
        raise ValueError("파일명이 올바르지 않습니다.")
    if filename != filename.replace("\\", "/").split("/")[-1]:
        raise ValueError("파일명에는 경로를 넣을 수 없습니다.")
    extension = Path(filename).suffix.lower()
    format_name = FORMATS.get(extension)
    if not format_name:
        raise ValueError("HTML, PDF, XLSX 파일만 업로드할 수 있습니다.")
    if not data or len(data) > MAX_FILE_BYTES:
        raise ValueError("생성 파일은 1바이트 이상 10MB 이하로 올려 주세요.")

    note = ""
    text_available = True
    extraction_limited = False
    if format_name == "html":
        try:
            extracted = data.decode("utf-8-sig")
        except UnicodeDecodeError as error:
            raise ValueError("HTML 파일은 UTF-8로 저장해 주세요.") from error
        if not extracted.strip():
            raise ValueError("HTML 파일이 비어 있습니다.")
    elif format_name == "pdf":
        if not data.startswith(b"%PDF-"):
            raise ValueError("PDF 파일 형식이 올바르지 않습니다.")
        try:
            from pypdf import PdfReader
            reader = PdfReader(io.BytesIO(data), strict=False)
            if reader.is_encrypted:
                raise ValueError("암호가 걸린 PDF는 평가할 수 없습니다.")
            if len(reader.pages) > 100:
                raise ValueError("PDF는 최대 100페이지까지 지원합니다.")
            chunks = []
            any_text = False
            for index, page in enumerate(reader.pages, 1):
                page_text = page.extract_text() or ""
                any_text = any_text or bool(page_text.strip())
                chunks.append(f"[PDF {index}페이지]\n{page_text}")
                if sum(map(len, chunks)) > MAX_TEXT_CHARS:
                    extraction_limited = index < len(reader.pages)
                    break
            extracted = "\n\n".join(chunks)
            if not any_text:
                text_available = False
                note = "텍스트를 추출하지 못했습니다. 스캔 PDF는 OCR이 필요하며 내용 평가는 보류됩니다."
        except ValueError:
            raise
        except Exception as error:
            raise ValueError("PDF 내용을 읽을 수 없습니다.") from error
    else:
        try:
            with zipfile.ZipFile(io.BytesIO(data)) as archive:
                members = archive.infolist()
                if len(members) > 2000 or sum(member.file_size for member in members) > 40 * 1024 * 1024:
                    raise ValueError("XLSX 내부 크기가 제한을 초과합니다.")
                if "xl/workbook.xml" not in archive.namelist():
                    raise ValueError("XLSX 파일 형식이 올바르지 않습니다.")
            from openpyxl import load_workbook
            workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True, keep_links=False)
            try:
                chunks = []
                cell_count = 0
                text_length = 0
                extraction_limited = len(workbook.worksheets) > 30
                for sheet in workbook.worksheets[:30]:
                    chunks.append(f"[시트: {sheet.title}]")
                    text_length += len(chunks[-1])
                    max_row = min(sheet.max_row or 10_000, 10_000)
                    max_col = min(sheet.max_column or 100, 100)
                    extraction_limited = extraction_limited or (sheet.max_row or 0) > max_row or (sheet.max_column or 0) > max_col
                    for row_number, row in enumerate(sheet.iter_rows(max_row=max_row, max_col=max_col), 1):
                        values = [f"{cell.coordinate}={cell.value}" for cell in row if cell.value is not None]
                        if values:
                            line = " | ".join(values)
                            chunks.append(line)
                            text_length += len(line)
                            cell_count += len(values)
                        if row_number >= 10000 or cell_count >= 3000 or text_length > MAX_TEXT_CHARS:
                            extraction_limited = extraction_limited or cell_count >= 3000
                            break
                    if cell_count >= 3000 or text_length > MAX_TEXT_CHARS:
                        break
                extracted = "\n".join(chunks)
                text_available = cell_count > 0
                note = "시트당 앞의 1만 행·100열에서 셀 값만 추출했습니다. 수식 실행·서식·차트·이미지는 검증하지 않습니다."
            finally:
                workbook.close()
        except ValueError:
            raise
        except Exception as error:
            raise ValueError("XLSX 내용을 읽을 수 없습니다.") from error

    truncated = extraction_limited or len(extracted) > MAX_TEXT_CHARS
    return {"filename": filename, "format": format_name, "size": len(data),
            "sha256": hashlib.sha256(data).hexdigest(), "extension": extension,
            "extractedText": extracted[:MAX_TEXT_CHARS], "textTruncated": truncated,
            "textAvailable": text_available, "extractionNote": note}
